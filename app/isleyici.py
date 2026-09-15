"""
İşleyici — ham belge verisini Fiş Aktarım Şablonu satırlarına çevirir.
Her satır sözlük (14 sütun karşılığı):
  fisno, fis_tarih(iso), fis_aciklama, hesap, evrak_no, evrak_tarih(iso),
  detay, borc, alacak, belge_turu(MF)

Şimdilik iskelet: banka/fatura/cek için temel akış. OCR/metin satırlarını
kural motoruyla eşleştirir, dengeli çift kayıt üretir. Kural dosyasındaki
banka hesabı / karşı hesap mantığı geliştirilecek.
"""
import io, re
from datetime import datetime
from app.kurallar import norm, dosya_banka_anahtari, banka_kisa_adi

# Gerçek işlem değil, hesap özeti/metadata satırı olduğu belli olan açıklamalar
# (ör. "Sicil: 26920050511927770340630") — bunlar tabloya karışırsa hem hesap
# eşleşmez hem de anlamsız devasa bir "tutar" üretebilir.
_METADATA_RE = re.compile(
    r"^(SICIL|IBAN|SUBE|VKN|TCKN|MUSTERI NO|HESAP NO|HESAP SAHIBI|"
    r"TARIH ARALIGI|BAKIYE|EK HESAP LIMITI|ORTAK HESAP|RUMUZ|VB MUS NO|"
    r"HESAP TURU|HESAP HAREKETLERI|SAYFA NO)\b"
)
_MAKUL_TUTAR_UST_SINIR = 1_000_000_000  # bu üstü gerçek bir banka hareketi değil, hatalı okunmuş sayıdır
_SAAT_ONEK_RE = re.compile(r"^\d{1,2}:\d{2}(:\d{2})?\s+")  # açıklama başındaki "12:31:21 " gibi saat
_ISLEMNO_RE = re.compile(r"\b\d{10,}\b")  # İşlem/referans no gibi uzun sayı dizileri açıklamaya karışmasın
_FOOTER_RE = re.compile(r"www\.|Sicil Numaras|Dekont yerine|Uyuşmazlık halinde", re.IGNORECASE)
_FOOTER_BASLANGIC_RE = re.compile(r"^\*{3,}")  # "***Dekont yerine..." gibi dipnotlar; maskeli kart no ("5472********1297") ile karışmasın diye SADECE satır başı


def _ddmmyyyy(iso: str) -> str:
    try:
        y, m, d = map(int, iso.split("-"))
        return f"{d:02d}.{m:02d}.{y}"
    except Exception:
        return iso or ""


# ----------------------------------------------------------------- tarih/sayı ayrıştırma
def _tarih_iso(s):
    if s is None:
        return ""
    if isinstance(s, datetime):
        return f"{s.year:04d}-{s.month:02d}-{s.day:02d}"
    s = str(s).strip()
    # Akbank internet şubesi CSV exportunda "Tarih" hücresi
    # "2026-08-31-11.28.21.190812" gibi ISO tarihin hemen ardına (boşluksuz,
    # tire ile) saat/mikrosaniye eklenmiş halde gelir — baştaki YYYY-MM-DD'yi
    # doğrudan al, geri kalanı (saat) yok say.
    m = re.match(r"^(\d{4}-\d{2}-\d{2})(?:[-T ].*)?$", s)
    if m:
        return m.group(1)
    # "20.08.2026 13:24" gibi tarih+saat birleşik hücrelerde (ör. Denizbank
    # Excel ekstresi) saat kısmını at, sadece tarihi ayrıştır.
    s = re.sub(r"\s+\d{1,2}:\d{2}(:\d{2})?\s*$", "", s)
    for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%y"):
        try:
            d = datetime.strptime(s, fmt)
            return f"{d.year:04d}-{d.month:02d}-{d.day:02d}"
        except Exception:
            pass
    return ""


def _sayi(s):
    if s is None or s == "":
        return None
    if isinstance(s, (int, float)):
        return float(s)
    s = str(s).strip().replace(" ", "")
    if "," in s and "." in s:
        # Hangisi ondalık ayraç belirsiz: TR "1.066,40" mı yoksa US "1,066.40" mı?
        # Sağdaki (sonda gelen) ayracı ondalık kabul et, diğerini binlik say.
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return float(s)
    except Exception:
        return None


_PARA_BIRIMI_RE = re.compile(r"\b(TRY|TL|USD|EUR|GBP|CHF|JPY)\b", re.IGNORECASE)


def _para_birimi_tespit(hucre) -> str:
    """Bir tutar hücresinin metninden para birimini çıkarır (ör. '-537,50 EUR'
    -> 'EUR'); bulunamazsa (bankaların çoğu TL'de hiç sütun eklemez) 'TL'
    varsayılır."""
    if hucre is None:
        return "TL"
    m = _PARA_BIRIMI_RE.search(str(hucre))
    if not m:
        return "TL"
    pb = m.group(1).upper()
    return "TL" if pb == "TRY" else pb


def _tarih_saat_dt(ham):
    """Akbank CSV'sindeki 'YYYY-MM-DD-HH.MM.SS.mikrosaniye' ham tarih hücresini
    saniye hassasiyetinde datetime'a çevirir (banka içi döviz dönüşüm
    çiftlerini eşleştirmek için) — ayrıştıramazsa None döner."""
    if ham is None:
        return None
    s = str(ham).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})-(\d{2})\.(\d{2})\.(\d{2})", s)
    if not m:
        return None
    y, mo, d, h, mi, se = map(int, m.groups())
    try:
        return datetime(y, mo, d, h, mi, se)
    except Exception:
        return None


# ----------------------------------------------------------------- fiş satırı
def _sat(fisno, tarih, aciklama, hesap, borc, alacak, evrak_no="", detay="", kaynak="", kaynak_dosya="", iade=False, sayfa=None,
          para_birimi="", kur=None, doviz_tutar=None):
    return {
        "fisno": fisno, "fis_tarih": tarih, "fis_aciklama": aciklama,
        "hesap": hesap, "evrak_no": evrak_no, "evrak_tarih": tarih,
        "detay": detay or aciklama, "borc": round(borc, 2), "alacak": round(alacak, 2),
        "belge_turu": "MF", "kaynak": kaynak, "kaynak_dosya": kaynak_dosya, "iade": iade,
        "sayfa": sayfa,  # kaynak PDF'te kaçıncı sayfa (fatura görseli için) — Excel kaynaklıysa None
        # Döviz cinsinden bir banka hareketi TL karşılığıyla işlendiğinde
        # (bkz. _banka_doviz_isle) bu üç alan doldurulur; TL işlemlerde boş kalır.
        "para_birimi": para_birimi, "kur": kur, "doviz_tutar": doviz_tutar,
    }


# ----------------------------------------------------------------- excel tablo -> satırlar
def _tablo_satirlari(hamlar):
    """Excel/PDF tablolarından (tarih, açıklama, tutar, borc/alacak) çıkarımı dener.
    Başlık satırını bulup sütunları eşler. Genel amaçlı, banka dökümü odaklı."""
    kayitlar = []
    for h in hamlar:
        for tablo in h.get("tablolar", []):
            if not tablo:
                continue
            bas_idx = None; harita = {}
            for i, row in enumerate(tablo[:30]):
                nrow = [norm(str(c)) for c in row]
                satir_harita = {}
                for j, c in enumerate(nrow):
                    if c in ("TARIH", "ISLEM TARIHI", "VALOR", "VALOR TARIHI"):
                        satir_harita["tarih"] = j
                    elif "ACIKLAMA" in c or "ISLEM" == c:
                        satir_harita["aciklama"] = j
                    elif ("TUTAR" in c or "MIKTAR" in c) and "BAKIYE" not in c:
                        # "TUTAR", "İşlem Tutarı", "Tutar (TL)" (norm -> "TUTAR TL") gibi
                        # varyasyonları yakalar; "Güncel Bakiye (TL)" gibi bakiye
                        # sütunlarını (norm -> "GUNCEL BAKIYE TL") hariç tutar.
                        satir_harita["tutar"] = j
                    elif "BORC" in c:
                        satir_harita["borc"] = j
                    elif "ALACAK" in c:
                        satir_harita["alacak"] = j
                    elif "GONDER" in c or "UNVAN" in c or "FIRMA" in c:
                        satir_harita.setdefault("aciklama", j)
                # Her aday satır BAĞIMSIZ değerlendirilir (önceki satırlardan
                # kalıntı eşleşme taşınmaz) — gerçek başlık satırı tek satırda
                # hem tarih hem tutar/borç/alacak sütununu birlikte içerir.
                if "tarih" in satir_harita and ("tutar" in satir_harita or "borc" in satir_harita or "alacak" in satir_harita):
                    harita = satir_harita
                    bas_idx = i; break
            if bas_idx is None:
                continue
            for row in tablo[bas_idx + 1:]:
                def g(key):
                    j = harita.get(key)
                    return row[j] if j is not None and j < len(row) else None
                tarih = _tarih_iso(g("tarih"))
                acik = _SAAT_ONEK_RE.sub("", str(g("aciklama") or "").strip())
                if "borc" in harita or "alacak" in harita:
                    borc = _sayi(g("borc")) or 0
                    alacak = _sayi(g("alacak")) or 0
                    tutar = borc - alacak
                    para_birimi = _para_birimi_tespit(g("borc")) or _para_birimi_tespit(g("alacak"))
                else:
                    tutar = _sayi(g("tutar")) or 0
                    para_birimi = _para_birimi_tespit(g("tutar"))
                if not tarih and not acik:
                    continue
                if tutar == 0:
                    continue
                if _METADATA_RE.match(norm(acik)):
                    continue
                if abs(tutar) > _MAKUL_TUTAR_UST_SINIR:
                    continue
                kayitlar.append({"tarih": tarih, "tarih_ham": g("tarih"), "aciklama": acik, "tutar": tutar,
                                 "para_birimi": para_birimi, "dosya": h.get("dosya", "")})
    return kayitlar


def _ham_metin_satirlari(hamlar):
    """OCR/PDF ham metninden (tarih ... açıklama ... tutar) desenli satırları çıkarır."""
    kayitlar = []
    # Bazı bankalar (ör. Halkbank hesap özeti) tarihi "13-08-2026" gibi tire
    # ile ayırır — nokta/slash'ın yanı sıra tireyi de kabul et.
    tarih_re = re.compile(r"\b(\d{1,2}[./\-]\d{1,2}[./\-]\d{2,4})\b")
    tutar_re = re.compile(r"-?\d{1,3}(?:[.\s]\d{3})*(?:[.,]\d{2})|-?\d+[.,]\d{2}")
    for h in hamlar:
        metin = h.get("ham_metin", "")
        if not metin:
            continue
        satirlar = [s.strip() for s in metin.splitlines()]
        n = len(satirlar)
        i = 0
        while i < n:
            hat = satirlar[i]
            i += 1
            if not hat:
                continue
            tm = tarih_re.search(hat)
            # Tarih satırın hemen başında değilse (küçük OCR gürültüsü hariç)
            # bu bir işlem satırı değil, bir üstbilgi/etiket satırıdır
            # (ör. "Tarih Aralığı : 01.08.2026 - 31.08.2026") — atla.
            if not tm or tm.start() > 3:
                continue
            # Tutarları SADECE tarihten SONRAKİ kısımda ara; aksi halde
            # "03.08.2026" gibi tarihin kendi ".08" parçası yanlışlıkla bir
            # tutar sanılabiliyordu.
            kalan = hat[tm.end():]
            tutarlar = tutar_re.findall(kalan)
            if not tutarlar:
                continue
            tarih = _tarih_iso(tm.group(1))
            # Ekstrede aynı satırda hem İŞLEM TUTARI hem de (o hareket sonrası)
            # BAKİYE gösteriliyorsa (ör. Vakıfbank: "... -5.000,00 183.681,28
            # FAST Anlık Ödeme") sıradaki İLK tutar gerçek işlem tutarıdır;
            # sonraki(ler) bakiyedir — SON değeri almak bakiyeyi tutar sanardı.
            tutar = _sayi(tutarlar[0])
            if tutar is None or tutar == 0:
                continue
            if abs(tutar) > _MAKUL_TUTAR_UST_SINIR:
                continue
            acik = _ISLEMNO_RE.sub("", kalan)
            for t in tutarlar:
                acik = acik.replace(t, "")
            acik = re.sub(r"\s+", " ", acik).strip(" \t-|")
            acik = _SAAT_ONEK_RE.sub("", acik)
            if _METADATA_RE.match(norm(acik)):
                continue
            # Birçok ekstrede kısa "İşlem Adı" (ör. "FAST Anlık Ödeme") tek
            # başına karşı tarafı ayırt etmeye yetmez; asıl bilgi bir sonraki
            # satır(lar)a taşar (ör. "... ERDİNÇ ÇÖL hesabına giden FAST
            # ödemesi"). Yeni bir işlem/tarih başlayana, dipnot/altbilgi
            # metnine (İBAN/şube/dekont uyarısı vb.) ya da anlamsız kısa bir
            # parçaya (taranmış logo/altbilgi gürültüsü) çatana kadar bu
            # devam satırlarını açıklamaya ekle.
            ek = []
            while i < n and len(ek) < 4:
                sonraki = _ISLEMNO_RE.sub("", satirlar[i])
                sonraki = re.sub(r"\s+", " ", sonraki).strip()
                if not sonraki:
                    i += 1
                    continue
                tm2 = tarih_re.search(sonraki)
                if tm2 and tm2.start() <= 3:
                    break
                if _FOOTER_RE.search(sonraki) or _FOOTER_BASLANGIC_RE.match(sonraki) or _METADATA_RE.match(norm(sonraki)):
                    break
                if len(sonraki) < 8 or " " not in sonraki:
                    break
                ek.append(sonraki)
                i += 1
            if ek:
                acik = re.sub(r"\s+", " ", acik + " " + " ".join(ek)).strip()
            kayitlar.append({"tarih": tarih, "aciklama": acik, "tutar": tutar,
                             "para_birimi": "TL", "dosya": h.get("dosya", "")})
    return kayitlar


def _kayitlar(hamlar):
    """Her belge için AYRI AYRI: önce tablo, tablo çıkmazsa ham metin (OCR)
    satırları. Dosya bazında karar vermek önemli — aynı işlemde tablo tabanlı
    bir dosya (ör. Excel/CSV ekstresi) ile tablo çıkarılamayan bir dosya
    (ör. metin katmanlı taranmış PDF, ya da tabloyu tek hücreye sıkıştıran
    Halkbank ekstresi gibi PDF'ler) karışık yüklenirse, TEK bir global
    "hepsi tablo / hepsi ham-metin" seçimi ikincisini sessizce sıfır kayıtla
    atlardı (tablo verenin varlığı yeter sayılır, ötekine hiç bakılmazdı)."""
    kayitlar = []
    for h in hamlar:
        k = _tablo_satirlari([h])
        if not k:
            k = _ham_metin_satirlari([h])
        kayitlar.extend(k)
    return kayitlar


# ----------------------------------------------------------------- döviz (banka içi dönüşüm + tek taraflı hareketler)
# "DÖVİZ SATIM/ALIM" gibi bankanın kendi TL<->döviz dönüşüm işlemi etiketleri
# (norm() sonrası Türkçe aksan zaten sadeleşmiş olur: DÖVİZ -> DOVIZ).
_DOVIZ_ICI_RE = re.compile(r"DOVIZ\s*(SATIM|ALIM|ALIS|SATIS)")
_ESLESME_TOLERANSI_SN = 5  # aynı dönüşümün TL ve döviz bacağı arasında beklenen en fazla zaman farkı


def _banka_doviz_isle(kayitlar, uyarilar=None):
    """Birden fazla para birimindeki banka dosyaları (ör. Akbank TL+USD+EUR+GBP)
    BİRLİKTE işlendiğinde uygulanan döviz muhasebe yöntemi:

    1) Banka İÇİ dönüşümler (TL<->döviz): TL dosyasındaki ve döviz dosyasındaki
       "DÖVİZ SATIM/ALIM" satırları aynı saniyeye yakın zaman damgasıyla
       eşleştirilir. TL bacağı bankanın GERÇEKLEŞEN TL tutarıyla işlenmeye
       devam eder (TCMB kuru KULLANILMAZ, hesaplama yapılmaz) — döviz bacağı
       ayrı bir fiş üretmez, sadece TL bacağın Döviz Tutar/Kur/Para Birimi
       bilgisini doldurmak için kullanılır.
    2) Tek taraflı döviz hareketleri (banka masrafı, döviz hesabından doğrudan
       çekiş, yurtdışı müşteri tahsilatı gibi TL karşılığı olmayan satırlar):
       aynı gün varsa o günün banka-içi dönüşümlerinden (ağırlıklı ortalama,
       TL tutarına göre ağırlıklandırılır) çıkarılan kur, yoksa o para birimi
       için en yakın günün kuru ile TL karşılığına çevrilir; orijinal döviz
       tutarı ve kullanılan kur ayrıca etiketlenir. Referans kur hiç
       bulunamazsa (o para biriminde hiç banka-içi dönüşüm yoksa) satır YANLIŞ
       bir TL tutarıyla işlenmez — atlanır ve uyarı listesine düşer.

    TL dışı para birimi hiç yoksa (olağan/çoğunluk durum) bu fonksiyon devre
    dışı kalır — kayıtlar aynen döner, tek para birimli hiçbir firmayı/akışı
    etkilemez."""
    if uyarilar is None:
        uyarilar = []
    para_birimleri = {(k.get("para_birimi") or "TL") for k in kayitlar}
    if para_birimleri <= {"TL"}:
        return kayitlar

    tl_kayitlar = [k for k in kayitlar if (k.get("para_birimi") or "TL") == "TL"]
    doviz_kayitlar = [k for k in kayitlar if (k.get("para_birimi") or "TL") != "TL"]

    # 1) banka-içi dönüşüm eşleştirme
    kullanilan = set()
    icdonusum_kurlari = {}  # (tarih, para_birimi) -> [(agirlik_tl, kur), ...]
    for tl_k in tl_kayitlar:
        if not _DOVIZ_ICI_RE.search(norm(tl_k.get("aciklama", ""))):
            continue
        tl_zaman = _tarih_saat_dt(tl_k.get("tarih_ham"))
        if not tl_zaman:
            continue
        en_yakin = None; en_yakin_fark = None
        for d_k in doviz_kayitlar:
            if id(d_k) in kullanilan:
                continue
            if not _DOVIZ_ICI_RE.search(norm(d_k.get("aciklama", ""))):
                continue
            d_zaman = _tarih_saat_dt(d_k.get("tarih_ham"))
            if not d_zaman:
                continue
            fark = abs((tl_zaman - d_zaman).total_seconds())
            if fark <= _ESLESME_TOLERANSI_SN and (en_yakin_fark is None or fark < en_yakin_fark):
                en_yakin = d_k; en_yakin_fark = fark
        if en_yakin is None:
            continue
        kullanilan.add(id(en_yakin))
        doviz_tutar = en_yakin["tutar"]
        pb = en_yakin.get("para_birimi")
        if doviz_tutar:
            kur = round(abs(tl_k["tutar"]) / abs(doviz_tutar), 4)
            tl_k["kur"] = kur
            icdonusum_kurlari.setdefault((tl_k.get("tarih"), pb), []).append((abs(tl_k["tutar"]), kur))
        tl_k["para_birimi"] = pb
        tl_k["doviz_tutar"] = doviz_tutar
        tl_k["_karsi_dosya"] = en_yakin.get("dosya", "")

    kalan_doviz = [k for k in doviz_kayitlar if id(k) not in kullanilan]

    # günlük ağırlıklı ortalama kur (TL tutarına göre ağırlıklı)
    gunluk_kur = {}
    for (gun, pb), liste in icdonusum_kurlari.items():
        toplam_agirlik = sum(a for a, _ in liste)
        if toplam_agirlik:
            gunluk_kur[(gun, pb)] = sum(a * k for a, k in liste) / toplam_agirlik

    # 2) tek taraflı döviz hareketleri -> TL karşılığına çevir
    sonuc = list(tl_kayitlar)
    for k in kalan_doviz:
        pb = k.get("para_birimi")
        gun = k.get("tarih")
        kur = gunluk_kur.get((gun, pb))
        if kur is None and gun:
            aday_gunler = sorted({g for (g, p) in gunluk_kur if p == pb and g})
            if aday_gunler:
                try:
                    hedef = datetime.fromisoformat(gun)
                    en_yakin_gun = min(aday_gunler, key=lambda g: abs((datetime.fromisoformat(g) - hedef).days))
                    kur = gunluk_kur.get((en_yakin_gun, pb))
                except Exception:
                    kur = None
        if kur:
            k["doviz_tutar"] = k["tutar"]
            k["kur"] = round(kur, 4)
            k["tutar"] = round(k["tutar"] * kur, 2)
            sonuc.append(k)
        else:
            uyarilar.append(f"{k.get('dosya','')}: {k.get('aciklama','')[:40]} ({pb} {k.get('tutar')}) — "
                            f"bu ay için hiç banka-içi dönüşüm kuru bulunamadı, işlenemedi (elle girilmeli)")
    return sonuc


# ----------------------------------------------------------------- BANKA
def isle_banka(hamlar, km, fis0, banka_hesap_kodu=""):
    """
    Banka dökümü -> fiş. Aynı banka hesabına ait, aynı tarihli hareketler TEK
    fişte toplanır (fiş no ve fiş tarihi ortak, her hareket kendi Borç/Alacak
    satır çiftiyle fişe eklenir); farklı tarih veya farklı banka hesabı yeni
    bir fiş no başlatır.
    Bankaya para GİRİŞİ (alacak, +): banka borç / karşı alacak
    Bankadan ÇIKIŞ (borç, -): karşı borç / banka alacak

    Banka hesabı DOSYA BAZINDA belirlenir:
      1) dosyadaki IBAN/hesap no daha önce öğrenilmiş bir hesaba bağlıysa -> otomatik
      2) değilse elle verilen banka_hesap_kodu / kural excel (vadesiz ana) / mizanda
         tek 102.x hesap -> varsayılan olarak TÜM dosyalara uygulanır
      3) hâlâ yoksa 102.01.001'e düşer ve uyarı verir
    Aynı işlemde tek bir YENİ (henüz öğrenilmemiş) IBAN/hesap no ile birlikte elle
    banka_hesap_kodu verilmişse, bu eşleşme kalıcı öğrenilir — sonraki aylarda bu
    dosya, diğer bankalarla karışık bir toplu işlemde bile otomatik doğru hesaba düşer.
    """
    uyarilar = []
    kayitlar = _kayitlar(hamlar)
    if not kayitlar:
        for h in hamlar:
            if h.get("ham_metin"):
                uyarilar.append(f"{h.get('dosya')}: tablo çıkarılamadı, ham metin var — elle düzenleme gerekebilir")
        return [], uyarilar

    # Birden fazla para birimi (ör. Akbank TL+USD+EUR+GBP) birlikte yüklendiyse
    # döviz muhasebe yöntemini uygula (bkz. _banka_doviz_isle) — tek para
    # birimliyse (olağan durum) hiçbir şey değişmeden aynı liste döner.
    kayitlar = _banka_doviz_isle(kayitlar, uyarilar)

    # 1) her dosyanın IBAN/hesap no anahtarını çıkar, öğrenilmiş eşleşmeye bak
    dosya_anahtar = {}
    dosya_hesap = {}
    for h in hamlar:
        fn = h.get("dosya", "")
        anahtar = dosya_banka_anahtari(h)
        dosya_anahtar[fn] = anahtar
        if anahtar and anahtar in km.banka_eslestirme:
            dosya_hesap[fn] = km.banka_eslestirme[anahtar]

    # 2) elle verilen kod + tek bir YENİ anahtar varsa -> kalıcı öğren
    manuel_hesap = (banka_hesap_kodu or "").strip()
    if manuel_hesap:
        yeni_anahtarlar = {a for a in dosya_anahtar.values() if a and a not in km.banka_eslestirme}
        if len(yeni_anahtarlar) == 1:
            km.banka_hesap_ogren(yeni_anahtarlar.pop(), manuel_hesap)
            for fn, a in dosya_anahtar.items():
                if a and a in km.banka_eslestirme:
                    dosya_hesap[fn] = km.banka_eslestirme[a]

    # 3) hiçbir şekilde belirlenemeyen dosyalar için varsayılan zincir
    varsayilan_hesap = manuel_hesap
    if not varsayilan_hesap:
        for hh in km.kural["hesaplar"]:
            if hh["kod"].startswith("102") and ("VADESIZ" in norm(hh.get("kullanim", "")) or "ANA" in norm(hh.get("kullanim", ""))):
                varsayilan_hesap = hh["kod"]; break
    if not varsayilan_hesap:
        for hh in km.kural["hesaplar"]:
            if hh["kod"].startswith("102"):
                varsayilan_hesap = hh["kod"]; break
    if not varsayilan_hesap:
        mizan_banka = [k for k, a in km.hesaplar if k.startswith("102")]
        if len(mizan_banka) == 1:
            varsayilan_hesap = mizan_banka[0]
    varsayilan_kullanildi = False
    if not varsayilan_hesap:
        varsayilan_hesap = "102.01.001"
        varsayilan_kullanildi = True

    for fn, kod in sorted(dosya_hesap.items()):
        uyarilar.append(f"{fn}: {kod} olarak otomatik tanındı (IBAN/hesap no eşleşmesi)")
    eksik_dosyalar = sorted(fn for fn in dosya_anahtar if fn not in dosya_hesap)
    if eksik_dosyalar and varsayilan_kullanildi:
        uyarilar.append("Banka hesabı belirlenemedi (kural/mizan'da net değil), 102.01.001 varsayıldı — "
                        "isteğe banka_hesap_kodu vererek düzeltebilirsiniz: " + ", ".join(eksik_dosyalar))

    fisler = []
    fis = fis0
    grup_fisno = {}  # (banka_hesap, tarih) -> fisno — aynı bankanın aynı tarihli hareketleri tek fişte toplanır
    for k in sorted(kayitlar, key=lambda x: (x["tarih"] or "", x.get("dosya", ""))):
        kd = k.get("dosya", "")
        banka_hesap = dosya_hesap.get(kd, varsayilan_hesap)
        grup_anahtar = (banka_hesap, k["tarih"])
        if grup_anahtar not in grup_fisno:
            grup_fisno[grup_anahtar] = f"{fis:05d}"
            fis += 1
        fisno = grup_fisno[grup_anahtar]
        karsi_dosya = k.get("_karsi_dosya")
        if karsi_dosya:
            # Banka içi TL<->döviz dönüşümünün karşı bacağı: km.eslestir()
            # açıklama metninden (ör. "DÖVİZ SATIM") hangi para biriminin
            # hesabı olduğunu ayırt edemez — karşı hesabı doğrudan eşleşen
            # döviz dosyasının KENDİ banka hesabı olarak al.
            karsi = dosya_hesap.get(karsi_dosya, varsayilan_hesap)
            kaynak = "banka"
        else:
            karsi, kaynak = km.eslestir(k["aciklama"])
            if not karsi:
                karsi = _CARI_BULUNAMADI_HESABI
                uyarilar.append(f"{k['aciklama'][:30]}: hesap eşleşmedi, {_CARI_BULUNAMADI_HESABI} varsayıldı")
        tutar = abs(k["tutar"])
        # Fiş Açıklama = "BANKA-TARİH" (ör. "YKB-05.04.2026"); Detay Açıklama
        # ise gerçek işlem metni. Tüm bankalarda aynı kural geçerli.
        fis_aciklama = f"{banka_kisa_adi(km.hesap_adi(banka_hesap), banka_hesap)}-{_ddmmyyyy(k['tarih'])}"
        aciklama = k["aciklama"]
        pb = k.get("para_birimi") or ""
        kur = k.get("kur")
        doviz_tutar = k.get("doviz_tutar")
        if k["tutar"] >= 0:
            fisler.append(_sat(fisno, k["tarih"], fis_aciklama, banka_hesap, tutar, 0, detay=aciklama, kaynak="banka", kaynak_dosya=kd,
                               para_birimi=pb, kur=kur, doviz_tutar=doviz_tutar))
            fisler.append(_sat(fisno, k["tarih"], fis_aciklama, karsi, 0, tutar, detay=aciklama, kaynak=kaynak, kaynak_dosya=kd,
                               para_birimi=pb, kur=kur, doviz_tutar=doviz_tutar))
        else:
            fisler.append(_sat(fisno, k["tarih"], fis_aciklama, karsi, tutar, 0, detay=aciklama, kaynak=kaynak, kaynak_dosya=kd,
                               para_birimi=pb, kur=kur, doviz_tutar=doviz_tutar))
            fisler.append(_sat(fisno, k["tarih"], fis_aciklama, banka_hesap, 0, tutar, detay=aciklama, kaynak="banka", kaynak_dosya=kd,
                               para_birimi=pb, kur=kur, doviz_tutar=doviz_tutar))
    return fisler, uyarilar


# ----------------------------------------------------------------- FATURA
_KDV_HESAP_KODU = {10: "191.01.010", 20: "191.01.020"}  # oran -> indirilecek KDV hesabı (mizanda farklıysa kural dosyasından düzeltilebilir)


def _efatura_listesi_satirlari(hamlar):
    """Bazı e-Fatura entegratörlerinin ('gelen fatura listesi' Excel raporu)
    formatını tanır: Fatura No, Fatura Tarihi, Gönderici Adı, KDV %10, KDV %20,
    KDV %10 Matrah, KDV %20 Matrah, Ek Vergiler gibi sütunlar. Genel
    _tablo_satirlari banka dökümü odaklı (TARİH/AÇIKLAMA/TUTAR) olduğu için bu
    formatı tanımıyor — ayrı bir sütun haritası gerekiyor. Bu formatta ürün/
    hizmet açıklaması YOK, sadece gönderici (tedarikçi) adı var."""
    kayitlar = []
    for h in hamlar:
        for tablo in h.get("tablolar", []):
            if not tablo:
                continue
            bas_idx = None; harita = {}
            for i, row in enumerate(tablo[:10]):
                nrow = [norm(str(c)) for c in row]
                if not any(c == "FATURA NO" or c.startswith("FATURA NUMARA") for c in nrow):
                    continue
                for j, c in enumerate(nrow):
                    if c == "FATURA NO" or c.startswith("FATURA NUMARA"):
                        harita["no"] = j
                    elif c == "FATURA TARIHI":
                        harita["tarih"] = j
                    elif ("GONDERICI" in c or "ALICI" in c or "MUSTERI" in c) and ("ADI" in c or "UNVAN" in c):
                        # "ADI"/"UNVAN" şartı, aynı raporda AYRICA bulunabilen
                        # "Müşteri Bayi Kodu"/"Müşteri Kodu" gibi kod
                        # sütunlarının (gerçek AREL satış listesinde olduğu
                        # gibi) asıl "Alıcı Adı" sütununu ezmesini önler —
                        # ilk eşleşen (soldaki) kazanır.
                        harita.setdefault("gonderici", j)
                    elif c == "ACIKLAMA":
                        # Bazı entegratör raporlarında ayrı bir "Gönderici/Alıcı
                        # Adı" sütunu yok, tedarikçi unvanı "Açıklama" sütununda
                        # verilir — daha iyi bir aday yoksa bunu kullan.
                        harita.setdefault("gonderici", j)
                    elif "MATRAH" in c and "10" in c:
                        harita["matrah10"] = j
                    elif "MATRAH" in c and "20" in c:
                        harita["matrah20"] = j
                    elif "KDV" in c and "10" in c:
                        harita["kdv10"] = j
                    elif "KDV" in c and "20" in c:
                        harita["kdv20"] = j
                    elif "EK VERGI" in c:
                        harita["ek"] = j
                bas_idx = i
                break
            if bas_idx is None or "gonderici" not in harita or "tarih" not in harita:
                continue
            for row in tablo[bas_idx + 1:]:
                def g(key):
                    j = harita.get(key)
                    return row[j] if j is not None and j < len(row) else None
                tarih = _tarih_iso(g("tarih"))
                gonderici = str(g("gonderici") or "").strip()
                if not tarih and not gonderici:
                    continue
                kalemler = []
                for oran, mkey, kkey in ((10, "matrah10", "kdv10"), (20, "matrah20", "kdv20")):
                    matrah = _sayi(g(mkey)) or 0
                    kdv = _sayi(g(kkey)) or 0
                    if matrah or kdv:
                        kalemler.append((oran, matrah, kdv))
                ek = _sayi(g("ek")) or 0
                if not kalemler and not ek:
                    continue
                kayitlar.append({
                    "tarih": tarih, "gonderici": gonderici, "kalemler": kalemler,
                    "ek_vergiler": ek, "fatura_no": str(g("no") or "").strip(),
                    "dosya": h.get("dosya", ""),
                })
    return kayitlar


# ------------------------------------------------------- PDF tekil e-Fatura (entegratör baskısı)
# Excel entegratör listesi gelmediğinde (veya sadece PDF fatura elimizdeyse):
# her tedarikçi KENDİ fatura şablonunu kullanıyor (eLogo/e-Fatura portal
# "DocumentPrint" çıktısı) — ürün/hizmet kalem tablosunun sütun düzeni
# tedarikçiden tedarikçiye tamamen farklı. Bu yüzden kalem bazlı ayrıştırma
# yerine TÜM şablonlarda ortak bulunan üst bilgi alanları hedeflenir:
# Fatura No/ID, Tarih, "Mal/Hizmet Toplam Tutarı", "Hesaplanan KDV(...)",
# "Ödenecek Tutar" / "Vergiler Dahil Toplam Tutar". Genelde bu alanlar sayfa
# sonundaki özet tablosunda temiz (etiket, tutar) satırları olarak bulunur;
# tablo çıkmazsa (bazı şablonlar sınırsız/borderless) ham metinden aynı
# etiketler regex ile aranır.
_PDF_FATURA_ID_RE = re.compile(r"Fatura\s*ID\s*:?\s*([A-Za-z0-9]+)", re.IGNORECASE)
_PDF_FATURA_NO_RE = re.compile(r"Fatura\s*(?:No|Numaras[ıi])\s*:?\s*([A-Za-z0-9\-]+)", re.IGNORECASE)
_PDF_FATURA_TARIHI_RE = re.compile(r"Fatura\s*Tarih[i]?\s*:?\s*(\d{1,2})[\s./-]+(\d{1,2})[\s./-]+(\d{2,4})", re.IGNORECASE)
_PDF_TARIH_FALLBACK_RE = re.compile(r"(?:^|\n)\s*Tarih\s*:?\s*(\d{1,2})[\s./-]+(\d{1,2})[\s./-]+(\d{2,4})", re.IGNORECASE | re.MULTILINE)
_PDF_TARIH_ANY_RE = re.compile(r"\b(\d{1,2})[.\-](\d{1,2})[.\-](\d{4})\b")
_PDF_IADE_RE = re.compile(r"Fatura\s*Tipi\s*:?\s*IADE", re.IGNORECASE)
_PDF_ODENECEK_RE = re.compile(r"(?:ÖDENECEK\s+(?:TUTAR|TOPLAM)|Ödenecek\s+Tutar)\s*:?\s*([\d.,]+)", re.IGNORECASE)
_PDF_MAL_HIZMET_RE = re.compile(r"Mal\s*/?\s*Hizmet\s*Toplam\s*Tutar[ıi]\s*:?\s*([\d.,]+)\s*(?:TL|TRY)?", re.IGNORECASE)
_PDF_HESAPLANAN_KDV_METIN_RE = re.compile(r"Hesaplanan\s+KDV[^(%\n]*[(%][^)]*?(\d+)[^)]*\)?\s*:?\s*([\d.,]+)\s*(?:TL|TRY)", re.IGNORECASE)
_PDF_KDV_ORAN_RE = re.compile(r"%\s*(\d+)")
_PDF_UNVAN_RE = re.compile(r"(ŞİRKETİ|SIRKETI|A\.\Ş\.|A\.S\.|LTD\.|ŞTİ\.|STI\.|ANONIM)", re.IGNORECASE)
_PDF_GONDERICI_KESME_RE = re.compile(r"\s{2,}|HİZMET\s*NO|HIZMET\s*NO|Tel:|Faks|VKN|Vergi", re.IGNORECASE)
_PDF_GONDERICI_DUR_RE = re.compile(
    r"VKN|TCKN|Vergi\s*Dairesi|Vergi\s*No|Tel:|Faks|Web\s*Sitesi|E-Posta|E-posta|"
    r"Mah\.|Mahallesi|\bMh\.|Sokak|Sok\.|\bSk\.|Cadde|Cad\.|\bCd\.|No\s*:\s*\d", re.IGNORECASE)
_PDF_METADATA_ALAN_RE = re.compile(r"Özelleştirme|Senaryo|Fatura\s*ID|Fatura\s*Tarih|Fatura\s*Tipi", re.IGNORECASE)
_PDF_TUTAR_HUCRE_RE = re.compile(r"^[\d.,]+\s*(?:TL|TRY)?$")
# SATIŞ faturalarında sayfanın üst kısmı KENDİ firmamızın (düzenleyenin)
# antetidir, karşı taraf (alıcı/müşteri) değil — bu yüzden ALIŞ'ta kullanılan
# "sayfanın ilk satırları" sezgisi satışta yanlış tarafı (kendi firmamızı)
# yakalar. Bunun yerine "SAYIN"/"ALICI" etiketi aranır — gerçek eLogo
# portal şablonunda bu etiket genelde KENDİ satırında yalnız durur, alıcı
# unvanı bir/iki satır ALTINDA gelir (aynı satırda değil).
_PDF_SAYIN_ETIKET_RE = re.compile(r"^\s*(?:SAY[İI]N|AL[İI]C[İI])\s*:?\s*(.*)$", re.IGNORECASE)
# Telekom tipi faturalarda (Turkcell/TTNET/Vodafone vb.) her vergi/kesinti
# satırı "<isim> %oran (Matrah:tutar [TRY|TL]) vergi_tutarı" biçimindedir —
# "Katma Deger Vergisi %20 (Matrah:544.23 TRY) 108.85" gibi.
_PDF_VERGI_SATIRI_RE = re.compile(
    r"^(.{3,60}?)\s*%\s*(\d+)\s*\(Matrah\w*\s*:?\s*([\d.,]+)\s*(?:TRY|TL)?\s*\)\s*([\d.,]+)\s*$",
    re.IGNORECASE | re.MULTILINE)


def _pdf_alici_adi_bul(metin: str) -> str:
    """SATIŞ faturasında 'SAYIN'/'ALICI' etiketinin altındaki (veya aynı
    satırındaki) alıcı unvanını çıkarır. Etiketten sonraki satırlarda
    aradaki kısa/anlamsız satırlar (logo/damga taşması gibi tek harfli
    satırlar) atlanır, adres/VKN gibi bilgiye (_PDF_GONDERICI_DUR_RE)
    çatılınca durulur; unvan 1-2 satıra yayılabilir."""
    satirlar_all = [s.strip() for s in metin.splitlines()]
    for i, s in enumerate(satirlar_all):
        m = _PDF_SAYIN_ETIKET_RE.match(s)
        if not m:
            continue
        aday = []
        inline = m.group(1).strip()
        if inline:
            aday.append(inline)
        for s2 in satirlar_all[i + 1:i + 12]:
            if len(aday) >= 2:
                break
            if not s2 or len(s2) < 3:
                continue
            if _PDF_GONDERICI_DUR_RE.search(s2):
                break
            aday.append(s2)
        if aday:
            return " ".join(aday)
    return ""


def _pdf_tekil_fatura_ayikla(ham, yon="alis"):
    """Tek bir PDF sayfasını (=tek fatura, tedarikçiye özgü şablon) ayrıştırır.
    _efatura_listesi_satirlari'nin döndürdüğüyle aynı sözlük şeklini üretir
    (tarih/gonderici/kalemler/ek_vergiler/fatura_no/dosya) + ayrıca 'iade'.
    yon='satis' iken 'gonderici' alanı aslında ALICI/müşteri adını taşır."""
    metin = ham.get("ham_metin", "")
    if not metin.strip() or len(metin) < 150:
        return None

    # "FATURA DETAYI" / İrsaliye manifest eki: bir e-faturanın (ör. kargo)
    # arkasına eklenen, "Gönderici Adı"/"Alıcı Adı" sütunlu sevkiyat tablosu —
    # kendi başına bir fatura değildir (aynı faturanın rakamlarını tekrarlar,
    # fatura no'su da yoktur), yoksa mükerrer kayıt üretilir.
    for tablo in ham.get("tablolar", []):
        for row in tablo[:3]:
            nrow = [norm(str(c)) for c in row]
            if any("GONDERICI ADI" in c for c in nrow) and any("ALICI ADI" in c for c in nrow):
                return None

    m = _PDF_FATURA_ID_RE.search(metin) or _PDF_FATURA_NO_RE.search(metin)
    fatura_no = m.group(1).strip() if m else ""

    m = _PDF_FATURA_TARIHI_RE.search(metin) or _PDF_TARIH_FALLBACK_RE.search(metin)
    d = mo = y = None
    if m:
        d, mo, y = m.groups()
    else:
        satirlar_ = metin.splitlines()
        for i, s in enumerate(satirlar_):
            if re.search(r"Fatura\s*Tarih", s, re.IGNORECASE):
                for s2 in satirlar_[i:i + 3]:
                    m2 = _PDF_TARIH_ANY_RE.search(s2)
                    if m2:
                        d, mo, y = m2.groups()
                        break
                break
    tarih = _tarih_iso(f"{d}.{mo}.{y}") if d else ""

    satirlar_all0 = [s.strip() for s in metin.splitlines()]
    satirlar = satirlar_all0[1:]  # ilk satır: portal/tarih üstbilgisi (konumsal sezgide)
    gonderici = ""
    if yon == "satis":
        # Önce "SAYIN"/"ALICI" etiketini dene — sayfanın üst kısmındaki eski
        # sezgi (aşağıda) satışta kendi firmamızı yakalar.
        gonderici = _pdf_alici_adi_bul(metin)
    if not gonderici:
        # 1) Öncelik: "ŞİRKETİ"/"A.Ş."/"LTD."/"ŞTİ."/"ANONIM" gibi unvan soneki —
        #    konumdan bağımsız en güvenilir sinyal. İki sütunlu şablonlarda
        #    (ör. "Turkcell Satış A.Ş. Arc Kimyevi ... Şirketi") düzenleyen ve
        #    müşteri unvanı aynı satıra sıkışabilir: yalnızca İLK soneke kadar
        #    kesilir (bitişik/ardışık sonekler — "LTD. ŞTİ." gibi — birleştirilir).
        #    "SAYIN" etiketine varılınca durulur (ondan sonrası alıcı bloğu).
        for s in satirlar_all0[:20]:
            if not s or s.lower() in ("e-fatura", "e-fatura."):
                continue
            if _PDF_SAYIN_ETIKET_RE.match(s):
                break
            if _PDF_METADATA_ALAN_RE.search(s):
                continue
            if re.match(r"^\s*(BANKA|ŞUBE|SUBE|HESAP|IBAN)\s*:", s, re.IGNORECASE):
                # "BANKA : GARANTI BANKASI A.S." gibi satırlar unvan soneki
                # ("A.S.") içerebilir ama bu banka adıdır, tedarikçi değil.
                continue
            eslesmeler = list(_PDF_UNVAN_RE.finditer(s))
            if not eslesmeler:
                continue
            bitis = eslesmeler[0].end()
            for sonraki in eslesmeler[1:]:
                if sonraki.start() - bitis <= 3:
                    bitis = sonraki.end()
                else:
                    break
            gonderici = s[:bitis].strip()
            break
    if not gonderici:
        # 2) Unvan soneki yoksa (şahıs/TCKN'li tedarikçi gibi) eski konumsal
        #    sezgi: ilk anlamlı 1-2 satırı al, adres/VKN gibi bilgiye çatılınca dur.
        aday = []
        for s in satirlar:
            if not s or s.lower() in ("e-fatura", "e-fatura."):
                continue
            if _PDF_GONDERICI_DUR_RE.search(s):
                break
            aday.append(s)
            if len(aday) >= 2:
                break
        gonderici = " ".join(aday)
        if _PDF_METADATA_ALAN_RE.search(gonderici):
            # iki sütunlu üstbilgi (ör. telekom faturası) ilk satırları yanlış yakaladı
            gonderici = ""

    iade = bool(_PDF_IADE_RE.search(metin))

    # sayfadaki tüm tablo satırlarından (etiket_hücresi, tutar_hücresi) çifti çıkar
    ozet = {}
    for tablo in ham.get("tablolar", []):
        for row in tablo:
            hucreler = [c for c in row if c not in (None, "")]
            if len(hucreler) < 2:
                continue
            son = " ".join(str(hucreler[-1]).split()).lstrip(":").strip()
            if not _PDF_TUTAR_HUCRE_RE.match(son):
                continue
            deger = _sayi(son)
            if deger is None:
                continue
            etiket = " ".join(str(hucreler[-2]).split())
            ozet[etiket] = deger

    matrah = None
    kdv_kalemleri = []  # [(oran, tutar)]
    ek_vergi = 0.0
    toplam = None
    for etiket, deger in ozet.items():
        e = norm(etiket).replace(" ", "")  # Türkçe aksan/boşluk duyarsız (ör. "DEĞER" == "DEGER")
        if e.startswith("MALHIZMETTOPLAMTUTARI") or e.startswith("VERGIHARICTUTAR"):
            matrah = deger
        elif e == "TOPLAM" and matrah is None:
            # Basit tek-kalemli şablonlarda (ör. Turkcell Satış superbox faturası)
            # "Mal Hizmet Toplam Tutarı" yok, sade "Toplam" satırı matrah'tır.
            matrah = deger
        elif e.startswith("HESAPLANANKDV") or (e.startswith("HESAPLANAN") and ("KDV" in e or "KATMADEGER" in e)):
            rm = _PDF_KDV_ORAN_RE.search(etiket)
            kdv_kalemleri.append((int(rm.group(1)) if rm else 20, deger))
        elif "KDV" in e:
            # Genel "%20.0 KDV" gibi kısa etiketler (tablo başlığı yok)
            rm = _PDF_KDV_ORAN_RE.search(etiket)
            if rm:
                kdv_kalemleri.append((int(rm.group(1)), deger))
            else:
                ek_vergi += deger
        elif e.startswith("HESAPLANAN"):
            ek_vergi += deger
        elif (e.startswith("VERGILERDAHILTOPLAMTUTAR") or e.startswith("ODENECEKTUTAR")
              or e.startswith("GENELTOPLAM") or e.startswith("FATURATUTARI")):
            toplam = deger

    # tablo eksik/yoksa (borderless şablon) ham metinden aynı alanları dene
    if matrah is None:
        m = _PDF_MAL_HIZMET_RE.search(metin)
        if m:
            matrah = _sayi(m.group(1))
    if not kdv_kalemleri:
        m = _PDF_HESAPLANAN_KDV_METIN_RE.search(metin)
        if m:
            kdv_kalemleri.append((int(m.group(1)), _sayi(m.group(2))))
    if toplam is None:
        m = _PDF_ODENECEK_RE.search(metin)
        if m:
            toplam = _sayi(m.group(1))

    # Telekom tipi fatura (Turkcell/TTNET/Vodafone vb.): "Katma Değer Vergisi %20
    # (Matrah:544.23 TRY) 108.85" / "Özel İletişim Vergisi %10 (Matrah:...) ..."
    # gibi satırlar ne tabloda ne de yukarıdaki genel etiketlerle yakalanır —
    # her "<isim> %oran (Matrah:tutar) vergi_tutarı" satırını tara: KDV/Katma
    # Değer içerenler kdv_kalemleri'ne, diğerleri (ÖİV, Telsiz taksiti vb.)
    # ek_vergi'ye eklenir.
    if matrah is None and not kdv_kalemleri:
        for lbl, oran_s, matrah_s, tutar_s in _PDF_VERGI_SATIRI_RE.findall(metin):
            oran_i = int(oran_s)
            m_i = _sayi(matrah_s)
            t_i = _sayi(tutar_s) or 0
            lbl_n = norm(lbl)
            if "KDV" in lbl_n or "KATMA DEGER" in lbl_n:
                kdv_kalemleri.append((oran_i, t_i))
                if m_i and (matrah is None or m_i > matrah):
                    matrah = m_i
            else:
                ek_vergi += t_i

    if matrah is None and toplam is not None:
        kdv_toplam = sum(t for _, t in kdv_kalemleri)
        matrah = round(toplam - kdv_toplam - ek_vergi, 2)

    # Son mutabakat: toplam biliniyorsa fiş her zaman ona denk gelsin — eksik
    # kalan (ör. "Tahsilatına Aracılık Edilen Ödemeleriniz" gibi ayrıştırılamayan
    # kalemler) sessizce kaybolmasın, farkı ek_vergi'ye ekle.
    if toplam is not None:
        kdv_toplam = sum(t for _, t in kdv_kalemleri)
        fark = round(toplam - ((matrah or 0) + kdv_toplam + ek_vergi), 2)
        if abs(fark) >= 0.01:
            if matrah is None:
                matrah = round(toplam - kdv_toplam - ek_vergi, 2)
            else:
                ek_vergi = round(ek_vergi + fark, 2)

    if not (fatura_no or gonderici) or (matrah is None and toplam is None):
        return None

    if kdv_kalemleri:
        kalemler = [(oran, matrah if i == 0 else 0, tutar) for i, (oran, tutar) in enumerate(kdv_kalemleri)]
    elif matrah:
        kalemler = [(0, matrah, 0)]
    else:
        kalemler = []
    return {
        "tarih": tarih, "gonderici": gonderici, "kalemler": kalemler,
        "ek_vergiler": ek_vergi, "fatura_no": fatura_no, "iade": iade,
        "dosya": ham.get("dosya", ""),
    }


def _pdf_fatura_satirlari(hamlar, yon="alis"):
    """Excel entegratör listesi formatı değil, tek tek fatura görüntüsü/baskısı
    olan PDF'ler (her SAYFA = bir fatura, tedarikçi başına farklı şablon).
    belge_oku.pdf_oku 'sayfalar' alanında sayfa bazlı (ham_metin, tablolar)
    verir — tek dosyanın düzleştirilmiş ham_metin'i tüm faturaları birbirine
    karıştıracağından mutlaka sayfa bazlı işlenir."""
    kayitlar = []
    for h in hamlar:
        if h.get("tur") not in ("pdf", "pdf_ocr"):
            continue
        sayfalar = h.get("sayfalar") or [{"ham_metin": h.get("ham_metin", ""), "tablolar": h.get("tablolar", [])}]
        for sayfa_no, sayfa in enumerate(sayfalar, start=1):
            veri = {"ham_metin": sayfa.get("ham_metin", ""), "tablolar": sayfa.get("tablolar", []),
                    "dosya": h.get("dosya", "")}
            k = _pdf_tekil_fatura_ayikla(veri, yon=yon)
            if k:
                k["sayfa"] = sayfa_no  # kaynak PDF'te kaçıncı sayfa (1'den) — fatura görseli göstermek için
                kayitlar.append(k)
    return kayitlar


_TELEKOM_RE = re.compile(r"\bTTNET\b|\bVODAFONE\b|\bTURKCELL\b|T[UÜ]RK\s*TELEKOM|SUPERONLINE", re.IGNORECASE)
_OIV_HESAP_KODU = "689.01.002"  # TTNET/Vodafone/Turkcell gibi telekom faturalarındaki ÖİV kırılımı
# (gerçek geçmiş fiş listesinden doğrulandı: "689.01.002 — ÖİV"). Firmanın
# mizanında bu kod farklıysa Düzenle modunda elle düzeltilebilir.
_CARI_BULUNAMADI_HESABI = "198.01.001"  # gönderici adına uyan bir cari hesap bulunamazsa

# SATIŞ (yön ters): CARİ = Alıcılar (12x) BORÇ, KDV = Hesaplanan KDV (391.x) ALACAK,
# GELİR (600 serisi) ALACAK — ALIŞ'ın tam aynası.
_KDV_HESAP_KODU_SATIS = {10: "391.01.010", 20: "391.01.020"}
_CARI_BULUNAMADI_HESABI_SATIS = "120.01.001"  # alıcı adına uyan bir cari hesap bulunamazsa


def _isle_efatura_listesi(kayitlar, km, fis0):
    """e-Fatura entegratör listesindeki her satırı fişe çevirir: her fatura
    için CARİ (tedarikçiye göre eşleştirilir) ve GİDER (cari hesabın kendisine
    göre öğrenilir — bkz. KuralMotoru.gider_hesabi) hesapları ayrı ayrı
    bulunur, KDV oranı başına ayrı gider+KDV satır çifti açılır. Ek Vergiler:
    telekom operatörlerinde (TTNET/Vodafone/Turkcell/Türk Telekom/Superonline)
    ÖİV olduğu bilindiğinden ayrı bir 689.01.002 satırı olarak kırılır; diğer
    tedarikçilerde (ör. faktoring/BSMV) hâlâ gider tutarına eklenir."""
    uyarilar = []
    fisler = []
    fis = fis0
    for k in sorted(kayitlar, key=lambda x: (x["tarih"] or "", x.get("fatura_no", ""))):
        fisno = f"{fis:05d}"
        gonderici = k["gonderici"]
        iade = bool(k.get("iade"))
        telekom = bool(_TELEKOM_RE.search(gonderici))
        cari, kaynak = km.eslestir(gonderici)
        if not cari:
            uyarilar.append(f"{gonderici[:40]}: cari hesabı eşleşmedi")
        gider_varsayilan = km.gider_hesabi(cari, gonderici)
        if not gider_varsayilan:
            uyarilar.append(f"{gonderici[:40]}: gider hesabı bilinmiyor (Fiş Listesi'nden öğretin)")
        if iade:
            uyarilar.append(f"{gonderici[:40]} ({k.get('fatura_no','')}): İADE faturası — borç/alacak yönünü kontrol edin")

        kalemler = k["kalemler"] or [(0, 0, 0)]  # sadece Ek Vergiler varsa (ör. faktoring/BSMV) oransız tek satır
        ek_vergi = k.get("ek_vergiler") or 0
        toplam = 0.0
        ilk = True
        for oran, matrah, kdv in kalemler:
            # Telekom faturalarında Ek Vergiler (ÖİV) ayrı 689 satırına gider,
            # gider matrahına eklenmez — diğerlerinde (faktoring/BSMV) eskisi
            # gibi ilk kaleme eklenir.
            gider_matrah = matrah + (ek_vergi if (ilk and not telekom) else 0)
            ilk = False
            if gider_matrah == 0 and kdv == 0:
                continue
            toplam += gider_matrah + kdv
            # Açıklama (detay) sadece gönderici/cari adı olsun — KDV oranı,
            # "Ek Vergiler dahil" gibi ekler yazılmaz (İADE zaten ayrı bir
            # rozetle gösteriliyor, metne tekrar eklenmez).
            fisler.append(_sat(fisno, k["tarih"], gonderici, gider_varsayilan, gider_matrah, 0,
                                evrak_no=k.get("fatura_no", ""), detay=gonderici, kaynak=kaynak, iade=iade,
                                kaynak_dosya=k.get("dosya", ""), sayfa=k.get("sayfa")))
            if kdv:
                kdv_hesap = _KDV_HESAP_KODU.get(oran, "191.01.020")
                fisler.append(_sat(fisno, k["tarih"], gonderici, kdv_hesap, kdv, 0,
                                    evrak_no=k.get("fatura_no", ""), detay=gonderici, kaynak="kdv", iade=iade,
                                    kaynak_dosya=k.get("dosya", ""), sayfa=k.get("sayfa")))
        if telekom and ek_vergi:
            toplam += ek_vergi
            fisler.append(_sat(fisno, k["tarih"], gonderici, _OIV_HESAP_KODU, ek_vergi, 0,
                                evrak_no=k.get("fatura_no", ""), detay=gonderici, kaynak="oiv", iade=iade,
                                kaynak_dosya=k.get("dosya", ""), sayfa=k.get("sayfa")))
        if toplam:
            fisler.append(_sat(fisno, k["tarih"], gonderici, cari or _CARI_BULUNAMADI_HESABI, 0, round(toplam, 2),
                                evrak_no=k.get("fatura_no", ""), detay=gonderici,
                                kaynak=kaynak, iade=iade,
                                kaynak_dosya=k.get("dosya", ""), sayfa=k.get("sayfa")))
        fis += 1
    return fisler, uyarilar


def _isle_efatura_listesi_satis(kayitlar, km, fis0):
    """SATIŞ faturaları için _isle_efatura_listesi'nin ayna (mirror) mantığı:
    CARİ (Alıcılar 12x) BORÇLANIR, KDV (391.x Hesaplanan KDV) ve GELİR (600
    serisi) ALACAKLANIR — yön ALIŞ'ın tam tersidir. 'gider_hesabi' metodu
    burada isim değişmeden aynen kullanılır: main.py bu tip için ayrı bir
    {tip}_gider_eslestirme.json dosyası verdiğinden aynı genel mekanizma
    cari->gelir öğrenimi için de çalışır."""
    uyarilar = []
    fisler = []
    fis = fis0
    for k in sorted(kayitlar, key=lambda x: (x["tarih"] or "", x.get("fatura_no", ""))):
        fisno = f"{fis:05d}"
        alici = k["gonderici"]
        iade = bool(k.get("iade"))
        cari, kaynak = km.eslestir(alici)
        if not cari:
            uyarilar.append(f"{alici[:40]}: cari hesabı eşleşmedi")
        gelir_varsayilan = km.gider_hesabi(cari, alici)
        if not gelir_varsayilan:
            uyarilar.append(f"{alici[:40]}: gelir hesabı bilinmiyor (Fiş Listesi'nden öğretin)")
        if iade:
            uyarilar.append(f"{alici[:40]} ({k.get('fatura_no','')}): İADE faturası — borç/alacak yönünü kontrol edin")

        kalemler = k["kalemler"] or [(0, 0, 0)]
        ek_vergi = k.get("ek_vergiler") or 0
        toplam = 0.0
        ilk = True
        for oran, matrah, kdv in kalemler:
            gelir_matrah = matrah + (ek_vergi if ilk else 0)
            ilk = False
            if gelir_matrah == 0 and kdv == 0:
                continue
            toplam += gelir_matrah + kdv
            fisler.append(_sat(fisno, k["tarih"], alici, gelir_varsayilan, 0, gelir_matrah,
                                evrak_no=k.get("fatura_no", ""), detay=alici, kaynak=kaynak, iade=iade,
                                kaynak_dosya=k.get("dosya", ""), sayfa=k.get("sayfa")))
            if kdv:
                kdv_hesap = _KDV_HESAP_KODU_SATIS.get(oran, "391.01.020")
                fisler.append(_sat(fisno, k["tarih"], alici, kdv_hesap, 0, kdv,
                                    evrak_no=k.get("fatura_no", ""), detay=alici, kaynak="kdv", iade=iade,
                                    kaynak_dosya=k.get("dosya", ""), sayfa=k.get("sayfa")))
        if toplam:
            fisler.append(_sat(fisno, k["tarih"], alici, cari or _CARI_BULUNAMADI_HESABI_SATIS, round(toplam, 2), 0,
                                evrak_no=k.get("fatura_no", ""), detay=alici,
                                kaynak=kaynak, iade=iade,
                                kaynak_dosya=k.get("dosya", ""), sayfa=k.get("sayfa")))
        fis += 1
    return fisler, uyarilar


def isle_fatura(hamlar, km, fis0, yon="alis"):
    # Önce bilinen e-Fatura entegratör listesi formatını dene (Fatura No,
    # Fatura Tarihi, Gönderici/Alıcı Adı, KDV %10/%20 sütunları) — bu formatta
    # ürün/hizmet açıklaması yok, gider/gelir hesabı cari hesaptan öğrenilir.
    # Tek tek fatura PDF'leri (her sayfa kendi şablonunda bir fatura) de aynı
    # işleme ile birleştirilir — bir yüklemede hem Excel liste hem PDF
    # fatura birlikte gelebilir, biri diğerini geçersiz kılmaz.
    excel_kayitlar = _efatura_listesi_satirlari(hamlar)
    pdf_kayitlar = _pdf_fatura_satirlari(hamlar, yon=yon)
    if excel_kayitlar and pdf_kayitlar:
        # AYNI faturalar hem Excel entegratör listesinde hem tekil PDF
        # baskısında birlikte gelebilir (ör. "gelen fatura listesi.xlsx" +
        # o listedeki faturaların tek tek PDF çıktısı) — bu durumda aynı
        # fatura_no iki kez fişlenmesin diye, Excel'de zaten bulunan
        # fatura_no'lu PDF kayıtları atlanır (Excel'in yapılandırılmış
        # KDV/matrah kırılımı daha güvenilir kabul edilir). Ama PDF
        # kaydındaki sayfa/dosya bilgisi kaybolmasın diye — aksi halde
        # fatura görseli butonu hiç görünmez — atlanan her PDF kaydının
        # sayfa/dosya bilgisi eşleşen Excel kaydına aktarılır.
        pdf_by_no = {}
        for k in pdf_kayitlar:
            no = k.get("fatura_no", "").strip().upper()
            if no:
                pdf_by_no.setdefault(no, k)
        for k in excel_kayitlar:
            no = k.get("fatura_no", "").strip().upper()
            eslesen = pdf_by_no.get(no)
            if eslesen and not k.get("sayfa"):
                k["sayfa"] = eslesen.get("sayfa")
                k["dosya"] = eslesen.get("dosya", k.get("dosya", ""))
        excel_no = {k.get("fatura_no", "").strip().upper() for k in excel_kayitlar if k.get("fatura_no")}
        pdf_kayitlar = [k for k in pdf_kayitlar if k.get("fatura_no", "").strip().upper() not in excel_no]
    e_kayitlar = excel_kayitlar + pdf_kayitlar
    if e_kayitlar:
        if yon == "satis":
            return _isle_efatura_listesi_satis(e_kayitlar, km, fis0)
        return _isle_efatura_listesi(e_kayitlar, km, fis0)

    # Aksi halde eski/genel akış: basit "tarih + açıklama + tutar" listesi
    # (tek bir Excel sütun düzeni varsayar, sabit %20 KDV ile).
    uyarilar = []
    kayitlar = _kayitlar(hamlar)
    if not kayitlar:
        for h in hamlar:
            if h.get("ham_metin"):
                uyarilar.append(f"{h.get('dosya')}: tablo çıkarılamadı (OCR ham metin) — elle düzenleme gerekebilir")
        return [], uyarilar
    fisler = []
    fis = fis0
    for k in sorted(kayitlar, key=lambda x: x["tarih"] or ""):
        fisno = f"{fis:05d}"
        cari, kaynak = km.eslestir(k["aciklama"])
        karsi = km.gider_hesabi(cari, k["aciklama"])
        toplam = abs(k["tutar"])
        matrah = round(toplam / 1.20, 2); kdv = round(toplam - matrah, 2)
        if not karsi:
            uyarilar.append(f"{k['aciklama'][:30]}: {'gelir' if yon=='satis' else 'gider'} hesabı eşleşmedi")
        if yon == "satis":
            fisler.append(_sat(fisno, k["tarih"], k["aciklama"], karsi, 0, matrah, detay=k["aciklama"], kaynak=kaynak))
            fisler.append(_sat(fisno, k["tarih"], k["aciklama"], "391.01.020", 0, kdv, detay=k["aciklama"], kaynak="kdv"))
            fisler.append(_sat(fisno, k["tarih"], k["aciklama"], cari or _CARI_BULUNAMADI_HESABI_SATIS, toplam, 0, detay=k["aciklama"], kaynak=kaynak))
        else:
            fisler.append(_sat(fisno, k["tarih"], k["aciklama"], karsi, matrah, 0, detay=k["aciklama"], kaynak=kaynak))
            fisler.append(_sat(fisno, k["tarih"], k["aciklama"], "191.01.020", kdv, 0, detay=k["aciklama"], kaynak="kdv"))
            fisler.append(_sat(fisno, k["tarih"], k["aciklama"], cari or _CARI_BULUNAMADI_HESABI, 0, toplam, detay=k["aciklama"], kaynak=kaynak))
        fis += 1
    return fisler, uyarilar


# ----------------------------------------------------------------- ÇEK
# ----------------------------------------------------------------- masraf icmali (personel gider fişi)
_KDV_ORAN_HESAP_MASRAF = {1: "191.01.001", 10: "191.01.010", 20: "191.01.020"}


def _kdv_oran_yuzde(carpan):
    """KDV çarpanından (matrah*çarpan=genel toplam, ör. 1.10 -> %10) yüzdeyi çıkarır."""
    if not carpan:
        return None
    yuzde = round((carpan - 1) * 100)
    return yuzde if yuzde > 0 else None


def _masraf_excel_satirlari(hamlar):
    """Personel masraf icmali Excel formatını tanır (ör. 'PERAKENDE SATIŞ
    VESİKALARI İLE TEVSİK EDİLEN GİDERLER İCMALİ'): SIRA NO, TARİH, NO, FİRMA,
    MATRAH, KDV ORANI, KDV TUTARI, GENEL TOPLAM, HESAP KODU, GİDER TÜRÜ
    sütunları. Gider hesabı satırda zaten verili olduğundan kural motoruna/
    cari eşleştirmeye gerek yoktur — doğrudan kullanılır."""
    kayitlar = []
    for h in hamlar:
        for tablo in h.get("tablolar", []):
            if not tablo:
                continue
            bas_idx = None; harita = {}
            for i, row in enumerate(tablo[:15]):
                nrow = [norm(str(c)) for c in row]
                if "SIRA NO" not in nrow or "HESAP KODU" not in nrow:
                    continue
                for j, c in enumerate(nrow):
                    if c == "SIRA NO":
                        harita["sira"] = j
                    elif c == "TARIH":
                        harita["tarih"] = j
                    elif c == "NO":
                        harita["no"] = j
                    elif c == "FIRMA":
                        harita["firma"] = j
                    elif c == "MATRAH":
                        harita["matrah"] = j
                    elif "KDV" in c and "ORAN" in c:
                        harita["kdv_oran"] = j
                    elif "KDV" in c and "TUTAR" in c:
                        harita["kdv_tutar"] = j
                    elif "GENEL TOPLAM" in c:
                        harita["genel"] = j
                    elif "HESAP KODU" in c:
                        harita["hesap"] = j
                    elif "GIDER TURU" in c:
                        harita["gider_turu"] = j
                bas_idx = i
                break
            if bas_idx is None:
                continue
            for row in tablo[bas_idx + 1:]:
                def g(key):
                    j = harita.get(key)
                    return row[j] if j is not None and j < len(row) else None
                if g("sira") in (None, ""):
                    continue
                hesap = str(g("hesap") or "").strip()
                firma = str(g("firma") or "").strip()
                genel = _sayi(g("genel"))
                if not hesap or not firma or genel is None:
                    continue
                kayitlar.append({
                    "tarih": _tarih_iso(g("tarih")),
                    "firma": firma,
                    "matrah": _sayi(g("matrah")) or 0,
                    "kdv_carpan": _sayi(g("kdv_oran")),
                    "kdv_tutar": _sayi(g("kdv_tutar")) or 0,
                    "genel": genel,
                    "hesap": hesap,
                    "gider_turu": str(g("gider_turu") or "").strip(),
                    "belge_no": str(g("no") or "").strip(),
                    "dosya": h.get("dosya", ""),
                })
    return kayitlar


def isle_masraf(hamlar, km, fis0, karsi_hesap_kodu=""):
    """Personel masraf icmali (perakende satış vesikaları) -> HER SATIR AYRI
    fiş: gider hesabı (dosyada verili HESAP KODU) borç + KDV oranına göre
    191.xx borç + karşı hesap (Kasa/Ortak — UI'dan elle verilir) alacak.
    Açıklama sadece firma adıdır. Cari eşleştirme motoru KULLANILMAZ çünkü
    gider hesabı zaten dosyada satır satır verili."""
    uyarilar = []
    kayitlar = _masraf_excel_satirlari(hamlar)
    if not kayitlar:
        uyarilar.append("Masraf icmali tablo olarak okunamadı (SIRA NO + HESAP KODU sütunlu Excel bekleniyor)")
        return [], uyarilar
    karsi = (karsi_hesap_kodu or "").strip() or _CARI_BULUNAMADI_HESABI
    fisler = []
    fis = fis0
    for k in kayitlar:
        fisno = f"{fis:05d}"
        firma = k["firma"]
        matrah = round(k["matrah"], 2)
        kdv = round(k["kdv_tutar"], 2)
        genel = round(k["genel"], 2)
        yuzde = _kdv_oran_yuzde(k["kdv_carpan"])
        kdv_hesap = _KDV_ORAN_HESAP_MASRAF.get(yuzde) if yuzde else None
        if kdv and not kdv_hesap:
            # Bilinmeyen/tanımsız orandaki KDV'yi sessizce atmak yerine gider
            # hesabına dahil eder — fiş her durumda dengeli çıkar.
            uyarilar.append(f"{firma}: KDV oranı tanımlı değil (%{yuzde}), KDV tutarı gider hesabına dahil edildi")
            matrah = round(matrah + kdv, 2)
            kdv = 0
        fisler.append(_sat(fisno, k["tarih"], firma, k["hesap"], matrah, 0,
                            evrak_no=k["belge_no"], detay=firma, kaynak="masraf",
                            kaynak_dosya=k["dosya"]))
        if kdv:
            fisler.append(_sat(fisno, k["tarih"], firma, kdv_hesap, kdv, 0,
                                evrak_no=k["belge_no"], detay=firma, kaynak="kdv",
                                kaynak_dosya=k["dosya"]))
        fisler.append(_sat(fisno, k["tarih"], firma, karsi, 0, genel,
                            evrak_no=k["belge_no"], detay=firma, kaynak="masraf",
                            kaynak_dosya=k["dosya"]))
        fis += 1
    return fisler, uyarilar


def isle_cek(hamlar, km, fis0):
    uyarilar = []
    kayitlar = _kayitlar(hamlar)
    if not kayitlar:
        uyarilar.append("Çek listesi tablo olarak okunamadı (Excel bekleniyor)")
        return [], uyarilar
    fisler = []
    fis = fis0
    for k in sorted(kayitlar, key=lambda x: x["tarih"] or ""):
        fisno = f"{fis:05d}"
        cari, kaynak = km.eslestir(k["aciklama"])
        tutar = abs(k["tutar"])
        fisler.append(_sat(fisno, k["tarih"], k["aciklama"], "101.01.001", tutar, 0, detay=k["aciklama"], kaynak="cek"))
        fisler.append(_sat(fisno, k["tarih"], k["aciklama"], cari or "120.01.001", 0, tutar, detay=k["aciklama"], kaynak=kaynak))
        fis += 1
    return fisler, uyarilar


def isle(tip, hamlar, km, fis0, banka_hesap_kodu=""):
    if tip == "banka":
        return isle_banka(hamlar, km, fis0, banka_hesap_kodu)
    if tip == "fatura":
        return isle_fatura(hamlar, km, fis0)
    if tip == "fatura_satis":
        return isle_fatura(hamlar, km, fis0, yon="satis")
    if tip == "masraf":
        return isle_masraf(hamlar, km, fis0, banka_hesap_kodu)
    if tip == "cek":
        return isle_cek(hamlar, km, fis0)
    return [], ["Bilinmeyen tip"]


# ----------------------------------------------------------------- çıktı (Fiş Aktarım Şablonu)
def fis_xlsx(satirlar):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Fiş Aktarım Şablonu"
    basliklar = ["Fiş No", "Fiş Tarihi", "Fiş Açıklama", "Hesap Kodu", "Evrak No",
                 "Evrak Tarihi", "Detay Açıklama", "Borç", "Alacak", "Miktar",
                 "Belge Türü", "Para Birimi", "Kur", "Döviz Tutar"]
    ws.append(basliklar)
    hf = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", start_color="1F4E78")
    thin = Side(style="thin", color="D0D0D0"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.font = hf; c.fill = hfill; c.alignment = Alignment(horizontal="center", vertical="center"); c.border = bd

    def dt(iso):
        try:
            y, m, d = map(int, iso.split("-")); return datetime(y, m, d)
        except Exception:
            return iso

    for r in satirlar:
        evno = r.get("evrak_no", "")
        evno = int(evno) if str(evno).isdigit() else evno
        ws.append([
            r.get("fisno", ""), dt(r.get("fis_tarih", "")), r.get("fis_aciklama", ""),
            r.get("hesap", ""), evno, dt(r.get("evrak_tarih", "")), r.get("detay", ""),
            r.get("borc") or None, r.get("alacak") or None, None,
            r.get("belge_turu", "MF"), r.get("para_birimi") or "", r.get("kur") or "",
            r.get("doviz_tutar") if r.get("doviz_tutar") is not None else "",
        ])
    sfill = PatternFill("solid", start_color="FFF3CD")
    for row in ws.iter_rows(min_row=2, max_row=len(satirlar) + 1):
        for c in row:
            c.border = bd
        row[1].number_format = "dd/mm/yyyy"; row[5].number_format = "dd/mm/yyyy"
        row[7].number_format = "#,##0.00"; row[8].number_format = "#,##0.00"
        if not row[3].value:
            row[3].fill = sfill
    for col, w in zip("ABCDEFGHIJKLMN", [10, 12, 26, 14, 20, 12, 30, 13, 13, 8, 10, 10, 8, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return bio


def fis_xml(satirlar, vir0=80001, tir0=950001):
    from collections import OrderedDict
    gruplar = OrderedDict()
    for r in satirlar:
        gruplar.setdefault(r.get("fisno", ""), []).append(r)
    now = datetime.now()
    def esc(s): return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    def ds(iso):
        try:
            y, m, d = map(int, iso.split("-")); return f"{d:02d}.{m:02d}.{y}"
        except Exception:
            return iso
    out = ['<?xml version="1.0" encoding="ISO-8859-9"?>', "<GL_VOUCHERS>"]
    vir = vir0; tir = tir0; ln = 0
    for fis, lines in gruplar.items():
        d = ds(lines[0].get("fis_tarih", "")); parcali = d.split(".")
        mm = parcali[1] if len(parcali) > 1 else "1"; yy = parcali[2] if len(parcali) > 2 else "2026"
        td = round(sum(l["borc"] for l in lines), 2); tc = round(sum(l["alacak"] for l in lines), 2)
        out += ['  <GL_VOUCHER DBOP="INS">',
                "    <INTERNAL_REFERENCE>%d</INTERNAL_REFERENCE>" % vir, "    <TYPE>4</TYPE>",
                "    <NUMBER>%s</NUMBER>" % fis, "    <DATE>%s</DATE>" % d,
                "    <TOTAL_DEBIT>%.2f</TOTAL_DEBIT>" % td, "    <TOTAL_CREDIT>%.2f</TOTAL_CREDIT>" % tc,
                "    <CREATED_BY>6</CREATED_BY>", "    <DATE_CREATED>%s</DATE_CREATED>" % now.strftime("%d.%m.%Y"),
                "    <HOUR_CREATED>%d</HOUR_CREATED>" % now.hour, "    <MIN_CREATED>%d</MIN_CREATED>" % now.minute,
                "    <SEC_CREATED>%d</SEC_CREATED>" % now.second, "    <CURRSEL_TOTALS>1</CURRSEL_TOTALS>",
                "    <DATA_REFERENCE>%d</DATA_REFERENCE>" % vir, "    <TRANSACTIONS>"]
        for l in lines:
            ln += 1; hes = l.get("hesap", ""); parent = hes.split(".")[0] if hes else ""
            out.append("      <TRANSACTION>")
            out.append("        <INTERNAL_REFERENCE>%d</INTERNAL_REFERENCE>" % tir)
            if l["alacak"] > 0: out.append("        <SIGN>1</SIGN>")
            out.append("        <GL_CODE>%s</GL_CODE>" % hes)
            out.append("        <PARENT_GLCODE>%s</PARENT_GLCODE>" % parent)
            if l["borc"] > 0: out.append("        <DEBIT>%.2f</DEBIT>" % l["borc"]); amt = l["borc"]
            else: out.append("        <CREDIT>%.2f</CREDIT>" % l["alacak"]); amt = l["alacak"]
            out += ["        <LINENO>%d</LINENO>" % ln, "        <DESCRIPTION>%s</DESCRIPTION>" % esc(l.get("detay", "")),
                    "        <TC_XRATE>1</TC_XRATE>", "        <TC_AMOUNT>%.2f</TC_AMOUNT>" % amt,
                    "        <QUANTITY>0</QUANTITY>", "        <DATA_REFERENCE>%d</DATA_REFERENCE>" % tir,
                    "        <DETLIST/>", "        <DEFNFLDLIST/>", "        <MONTH>%s</MONTH>" % int(mm),
                    "        <YEAR>%s</YEAR>" % int(yy), "        <DOC_DATE>%s</DOC_DATE>" % d,
                    "        <GUID>00000000-0000-0000-0000-%012d</GUID>" % tir, "      </TRANSACTION>"]
            tir += 1
        out += ["    </TRANSACTIONS>", "  </GL_VOUCHER>"]; vir += 1
    out.append("</GL_VOUCHERS>")
    return "\n".join(out).encode("iso-8859-9", "replace")
