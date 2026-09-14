"""
Belge okuyucu — gelen dosyayı ham satır/tablo verisine çevirir.
Desteklenen: .xlsx/.xls (openpyxl + eski .xls için xlrd), .pdf (pdfplumber; metin yoksa OCR),
             .png/.jpg/.jpeg (OCR: pytesseract), .csv (bankaların internet şubesi exportu).

Çıktı: {"tur": "excel|pdf|pdf_ocr|resim_ocr|csv", "tablolar": [[[hücre,...],...]],
         "ham_metin": "...", "uyari": "..."}
Not: OCR sonuçları hatalı olabilir; kullanıcı önizlemede düzeltir.
"""
import csv as _csv
import io
from pathlib import Path


def _ext(name: str) -> str:
    return Path(name).suffix.lower()


def _xls_oku(path: Path):
    """Eski .xls formatı için xlrd ile okuma."""
    import xlrd
    wb = xlrd.open_workbook(str(path))
    tablolar = []
    for sn in wb.sheet_names():
        ws = wb.sheet_by_name(sn)
        satirlar = []
        for r in range(ws.nrows):
            row = ws.row_values(r)
            if any(c not in (None, "") for c in row):
                satirlar.append([("" if c is None else c) for c in row])
        if satirlar:
            tablolar.append(satirlar)
    return {"tur": "excel", "tablolar": tablolar, "ham_metin": "", "uyari": ""}


def excel_oku(path: Path, ext: str = ""):
    ext = (ext or _ext(path.name)).lower()
    if ext == ".xls":
        try:
            return _xls_oku(path)
        except Exception as e:
            return {"tur": "excel", "tablolar": [], "ham_metin": "",
                    "uyari": f".xls dosyası okunamadı: {e}"}

    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        # .xlsx uzantılı ama içerik aslında eski .xls olabilir -> fallback dene
        try:
            return _xls_oku(path)
        except Exception as e:
            return {"tur": "excel", "tablolar": [], "ham_metin": "",
                    "uyari": f"Excel dosyası okunamadı (bozuk/desteklenmeyen format): {e}"}

    tablolar = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        satirlar = []
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                satirlar.append([("" if c is None else c) for c in row])
        if satirlar:
            tablolar.append(satirlar)
    return {"tur": "excel", "tablolar": tablolar, "ham_metin": "", "uyari": ""}


def csv_oku(path: Path):
    """Banka internet şubesi CSV exportu (ör. Akbank). Türk bankaları çoğunlukla
    eski Windows-1254 (Turkish) kod sayfasıyla ';' ayraçlı csv üretir — bunu
    UTF-8 sanıp okumaya çalışmak Türkçe karakterleri (İ,ğ,ş,ç,ö,ü) bozar
    (mojibake). Önce UTF-8 dener, olmazsa cp1254'e (sonra iso-8859-9'a) düşer."""
    data = path.read_bytes()
    if data.startswith(b"\xef\xbb\xbf"):
        data = data[3:]
    metin = None
    for enc in ("utf-8", "cp1254", "iso-8859-9"):
        try:
            metin = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if metin is None:
        metin = data.decode("utf-8", errors="replace")

    satirlar_ham = metin.splitlines()
    # Ayracı otomatik tespit et — TR banka exportlarında genelde ';' (tutarlarda
    # ondalık ayracı ',' olduğundan ',' güvenilir bir ayraç değildir).
    ornek = "\n".join(satirlar_ham[:8])
    try:
        ayrac = _csv.Sniffer().sniff(ornek, delimiters=";,\t").delimiter
    except Exception:
        ayrac = ";" if ornek.count(";") >= ornek.count(",") else ","

    tablo = []
    try:
        for row in _csv.reader(satirlar_ham, delimiter=ayrac):
            if any((c or "").strip() for c in row):
                tablo.append(row)
    except Exception as e:
        return {"tur": "csv", "tablolar": [], "ham_metin": metin,
                "uyari": f"CSV ayrıştırılamadı: {e}"}

    tablolar = [tablo] if tablo else []
    return {"tur": "csv", "tablolar": tablolar, "ham_metin": metin, "uyari": ""}


def pdf_oku(path: Path):
    """Önce metin/tablo dener; metin çıkmazsa OCR'a düşer.
    'sayfalar': sayfa bazlı (ham_metin, tablolar) listesi de eklenir — banka/çek
    modülleri hâlâ düzleştirilmiş ham_metin/tablolar kullanır (aynı davranış),
    ama tek sayfa = tek fatura gibi sayfa sınırının önemli olduğu akışlar
    (bkz. isleyici._pdf_fatura_satirlari) 'sayfalar'ı kullanabilir."""
    import pdfplumber
    tablolar = []
    ham = []
    sayfalar = []
    metin_var = False
    with pdfplumber.open(path) as pdf:
        for sayfa in pdf.pages:
            t = sayfa.extract_text() or ""
            sayfa_tablolar = []
            for tb in (sayfa.extract_tables() or []):
                temiz = [[("" if c is None else str(c)) for c in row] for row in tb]
                if temiz:
                    sayfa_tablolar.append(temiz)
                    tablolar.append(temiz)
            if t.strip():
                metin_var = True
                ham.append(t)
            sayfalar.append({"ham_metin": t, "tablolar": sayfa_tablolar})
    if metin_var or tablolar:
        return {"tur": "pdf", "tablolar": tablolar, "ham_metin": "\n".join(ham), "uyari": "",
                "sayfalar": sayfalar}
    return pdf_ocr(path)


def pdf_ocr(path: Path):
    """Taranmış PDF: her sayfayı görüntüye çevirip OCR."""
    try:
        import pdfplumber
        import pytesseract
        from PIL import Image
    except ImportError as e:
        return {"tur": "pdf_ocr", "tablolar": [], "ham_metin": "",
                "uyari": f"OCR kütüphanesi yok: {e}"}
    ham = []
    sayfalar = []
    with pdfplumber.open(path) as pdf:
        for sayfa in pdf.pages:
            im = sayfa.to_image(resolution=300).original
            txt = pytesseract.image_to_string(im, lang=_ocr_lang())
            ham.append(txt)
            sayfalar.append({"ham_metin": txt, "tablolar": []})
    return {"tur": "pdf_ocr", "tablolar": [], "ham_metin": "\n".join(ham),
            "uyari": "Taranmış PDF OCR ile okundu; verileri kontrol edin.",
            "sayfalar": sayfalar}


def resim_oku(path: Path):
    """Fotoğraf/tarama: OCR."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError as e:
        return {"tur": "resim_ocr", "tablolar": [], "ham_metin": "",
                "uyari": f"OCR kütüphanesi yok: {e}"}
    im = Image.open(path)
    try:
        im = im.convert("L")
    except Exception:
        pass
    txt = pytesseract.image_to_string(im, lang=_ocr_lang())
    return {"tur": "resim_ocr", "tablolar": [], "ham_metin": txt,
            "uyari": "Görüntü OCR ile okundu; verileri kontrol edin."}


def _ocr_lang():
    """Türkçe varsa tur+eng, yoksa eng."""
    try:
        import pytesseract
        langs = pytesseract.get_languages()
        if "tur" in langs:
            return "tur+eng"
    except Exception:
        pass
    return "eng"


def belge_oku(path: Path, orijinal_ad: str = ""):
    ext = _ext(orijinal_ad or path.name)
    if ext in (".xlsx", ".xls", ".xlsm"):
        return excel_oku(path, ext)
    if ext == ".pdf":
        return pdf_oku(path)
    if ext == ".csv":
        return csv_oku(path)
    if ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"):
        return resim_oku(path)
    return {"tur": "bilinmeyen", "tablolar": [], "ham_metin": "",
            "uyari": f"Desteklenmeyen dosya türü: {ext}"}
